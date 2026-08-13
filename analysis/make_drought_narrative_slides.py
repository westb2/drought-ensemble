#!/usr/bin/env python3
"""Build a short narrative slide deck from drought analysis figures."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

FIG = Path("/glade/derecho/scratch/bwest/drought-ensemble/analysis/figures")
OUT = FIG / "drought_narrative_slides.pptx"

# 16:9
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.4)
TITLE_H = Inches(0.55)

ACCENT = RGBColor(0x4A, 0x24, 0x10)
MUTED = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SECTION_BG = RGBColor(0x4A, 0x24, 0x10)


def _set_run(run, text, *, size=18, bold=False, color=None):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    if color is not None:
        run.font.color.rgb = color


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def add_title_bar(slide, text: str):
    box = slide.shapes.add_textbox(MARGIN, Inches(0.2), W - 2 * MARGIN, TITLE_H)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    _set_run(p.add_run(), text, size=22, bold=True, color=ACCENT)


def add_picture_fit(slide, path: Path, top=Inches(0.85), bottom_margin=Inches(0.35)):
    """Place image maximized in remaining slide area, centered."""
    if not path.exists():
        box = slide.shapes.add_textbox(MARGIN, top, W - 2 * MARGIN, Inches(1))
        _set_run(box.text_frame.paragraphs[0].add_run(), f"Missing: {path.name}", size=14, color=MUTED)
        return
    max_w = W - 2 * MARGIN
    max_h = H - top - bottom_margin
    # Let pptx scale; compute aspect from file via pillow if available
    try:
        from PIL import Image

        with Image.open(path) as im:
            iw, ih = im.size
        aspect = iw / ih
        # fit
        if max_w / max_h > aspect:
            h = max_h
            w = h * aspect
        else:
            w = max_w
            h = w / aspect
    except Exception:
        w, h = max_w, max_h
    left = (W - w) / 2
    slide.shapes.add_picture(str(path), left, top, width=w, height=h)


def slide_figure(prs, title: str, filename: str):
    slide = blank_slide(prs)
    add_title_bar(slide, title)
    add_picture_fit(slide, FIG / filename)
    return slide


def slide_section(prs, section: str, subtitle: str = ""):
    slide = blank_slide(prs)
    # full-bleed dark rectangle
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        W,
        H,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = SECTION_BG
    shape.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), W - Inches(1.6), Inches(1.2))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), section, size=36, bold=True, color=WHITE)

    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.9), W - Inches(1.6), Inches(0.8))
        tf2 = box2.text_frame
        tf2.clear()
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        _set_run(p2.add_run(), subtitle, size=18, color=RGBColor(0xE8, 0xC3, 0x9E))
    return slide


def slide_title(prs):
    slide = blank_slide(prs)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), W - Inches(1.6), Inches(1.4))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), "Drought storage memory", size=36, bold=True, color=ACCENT)

    box2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), W - Inches(1.6), Inches(1.2))
    tf2 = box2.text_frame
    tf2.clear()
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    _set_run(
        p2.add_run(),
        "Temporary vs persistent deficits · overland connectivity · local Budyko · near-surface shielding",
        size=16,
        color=MUTED,
    )
    p3 = tf2.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    _set_run(p3.add_run(), "Potomac & Wolf", size=14, color=MUTED)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    slide_title(prs)

    # ----- 1. Results -----
    slide_section(prs, "1. Results", "What temporary vs persistent looks like")
    slide_figure(
        prs,
        "Definition: temporary vs persistent recovery",
        "drought_recovery_temp_persist_definition.png",
    )
    slide_figure(prs, "Drought course (10-year)", "ten_year_drought_streamflow_storage.png")
    slide_figure(prs, "Drought course (50-year)", "fifty_year_drought_streamflow_storage.png")
    slide_figure(prs, "Recovery by drought length", "drought_recovery_totals_and_anomalies.png")
    slide_figure(
        prs,
        "Fraction of storage deficit remaining",
        "drought_recovery_by_length_fractional_storage.png",
    )
    slide_figure(
        prs,
        "Temporary saturates; persistent grows",
        "drought_recovery_temp_persistent_deficit_bars.png",
    )
    slide_figure(
        prs,
        "Temporary vs persistent storage — Potomac",
        "drought_recovery_temp_persistent_storage_deficit_maps_potomac2.png",
    )
    slide_figure(
        prs,
        "Temporary vs persistent storage — Wolf",
        "drought_recovery_temp_persistent_storage_deficit_maps_wolf2.png",
    )
    slide_figure(
        prs,
        "WTD anomaly through recovery — Potomac",
        "drought_recovery_wtd_anomaly_maps_potomac2.png",
    )
    slide_figure(
        prs,
        "WTD anomaly through recovery — Wolf",
        "drought_recovery_wtd_anomaly_maps_wolf2.png",
    )

    # ----- 2. Overland / drainage as predictor (recovery covariate group) -----
    slide_section(
        prs,
        "2. Overland flow → deficit type",
        "Low-overland cells hold the persistent mass; distance-to-stream is weaker",
    )
    slide_figure(
        prs,
        "Temporary / persistent deficits vs covariates",
        "drought_recovery_deficit_covariate_spearman.png",
    )
    slide_figure(
        prs,
        "Persistent mass concentrates in low-overland cells",
        "drought_recovery_drainage_deficit_concentration.png",
    )
    slide_figure(
        prs,
        "log overland flow vs distance to stream",
        "drought_recovery_drainage_threshold_bins.png",
    )
    slide_figure(
        prs,
        "Enrichment of top-20% persist below flow thresholds",
        "drought_recovery_drainage_threshold_enrichment.png",
    )
    slide_figure(
        prs,
        "Drainage-position heuristics vs temp/persist",
        "drought_recovery_drainage_position_heuristics.png",
    )
    slide_figure(
        prs,
        "Best single-factor maps (WTD, K, recharge, porosity)",
        "drought_recovery_best_factors_maps.png",
    )
    slide_figure(
        prs,
        "Best single-factor hexbins",
        "drought_recovery_best_factors_hexbin.png",
    )

    # ----- 3. Local Budyko framing -----
    slide_section(
        prs,
        "3. Local Budyko framing",
        "Domain climate sets the regime; within-basin ranking is overland / WTD, not φ",
    )
    slide_figure(
        prs,
        "Aridity is flat; overland participation varies",
        "budyko_phi_vs_overland_hist.png",
    )
    slide_figure(
        prs,
        "Maps: aridity, overland, temporary fraction",
        "budyko_maps_phi_overland_ftemp.png",
    )
    slide_figure(
        prs,
        "f_temp vs overland, WTD, and stream distance",
        "budyko_hexbin_ftemp_vs_predictors.png",
    )
    slide_figure(
        prs,
        "Local Budyko space colored by temporary fraction",
        "budyko_space_ftemp.png",
    )
    slide_figure(
        prs,
        "Temporary fraction by landscape stratum",
        "budyko_ftemp_by_stratum.png",
    )
    slide_figure(
        prs,
        "Temp / persist volume by stratum",
        "budyko_stratified_temp_persist_bars.png",
    )
    slide_figure(
        prs,
        "Within-domain Spearman (overland & WTD vs f_temp)",
        "budyko_spearman_ftemp.png",
    )

    # ----- 4. Mechanism (depth) -----
    slide_section(
        prs,
        "4. Mechanism: depth",
        "Near-surface regenerates; deep losses persist",
    )
    slide_figure(
        prs,
        "Temporary ≈ near-surface; persistent ≈ deep",
        "shallow_deep_temp_persist_partition.png",
    )
    slide_figure(
        prs,
        "Recovery fraction by depth (10-year)",
        "shallow_deep_layer_recovery_profile_10yr.png",
    )
    slide_figure(
        prs,
        "Spatial: near-surface fraction ≈ temporary fraction",
        "shallow_deep_spatial_10yr.png",
    )
    slide_figure(
        prs,
        "Mid-drought: near-surface regenerates, deep ratchets down",
        "shallow_deep_mid_drought_regen.png",
    )
    slide_figure(
        prs,
        "Same series through recovery (near-surface rebounds; deep lags)",
        "shallow_deep_mid_drought_regen_with_recovery.png",
    )
    slide_figure(
        prs,
        "One drought year: storage tracks precip only near the surface",
        "shallow_deep_mid_drought_year_pulse.png",
    )

    prs.save(OUT)
    print("wrote", OUT, f"({OUT.stat().st_size/1024:.0f} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
