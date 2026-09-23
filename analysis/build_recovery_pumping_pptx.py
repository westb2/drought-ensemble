#!/usr/bin/env python3
"""Build a Google-Slides-friendly PowerPoint of drought + pumping recovery figures."""
from __future__ import annotations

from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path("/glade/derecho/scratch/bwest/drought-ensemble")
sys.path.insert(0, str(ROOT))

from analysis.figure_paths import FIG_ROOT, fig_path, resolve_figure  # noqa: E402

FIG_DIR = FIG_ROOT
OUT = ROOT / "analysis" / "recovery_and_pumping_figures.pptx"

# (section title, slides: (title, filename, notes))
SECTIONS = [
    (
        "Drought recovery (1 / 3 / 10 / 50-year)",
        [
            (
                "Recovery: storage & flow by drought length",
                "drought_recovery_totals_and_anomalies.png",
                "5-year recovery after drought end. Baseline = continued average.",
            ),
            (
                "Fractional storage deficit during recovery",
                "drought_recovery_by_length_fractional_storage.png",
                "Fraction of initial (t=0) storage deficit remaining.",
            ),
            (
                "10-year drought course",
                "ten_year_drought_streamflow_storage.png",
                "Last 3 spinup years → drought → 5-year recovery.",
            ),
            (
                "50-year drought course",
                "fifty_year_drought_streamflow_storage.png",
                "Same layout for the 50-year member.",
            ),
            (
                "WTD anomaly maps (recovery)",
                "drought_recovery_wtd_anomaly_maps.png",
                "ΔWTD at recovery start, +1 yr, +5 yr. Positive = deeper / drier.",
            ),
            (
                "Temporary vs persistent storage deficits",
                "drought_recovery_temp_persistent_storage_deficit_maps.png",
                "Temporary = recovered in recovery year 1; persistent = remaining.",
            ),
            (
                "Temporary vs persistent WTD deficits",
                "drought_recovery_temp_persistent_wtd_deficit_maps.png",
                "Same split for ΔWTD.",
            ),
            (
                "Domain-total temporary vs persistent",
                "drought_recovery_temp_persistent_deficit_bars.png",
                "Map-based domain totals (10⁶ m³).",
            ),
            (
                "Fast spike vs persistent (corrected)",
                "drought_recovery_temp_persistent_deficit_bars_spike.png",
                "Slow branch projected to t=0 to isolate the fast spike.",
            ),
        ],
    ),
    (
        "Pumping analogues (3-year pumping tests)",
        [
            (
                "Design note",
                None,
                "Ensemble is 40 yr average spinup + 3 yr pumping at rates "
                "1e-7 … 1e-4 m/h (domain-average). No post-pump recovery years, "
                "so figures compare rates during pumping and split early (yr1) "
                "vs buildup (yr2–3) instead of temporary/persistent recovery.",
            ),
            (
                "During pumping: storage & flow by rate",
                "pumping_during_totals_and_anomalies.png",
                "t=0 at pumping onset; shaded = pumping period.",
            ),
            (
                "Fractional deficit growth during pumping",
                "pumping_by_rate_fractional_storage.png",
                "Deficit relative to end of pump year 1.",
            ),
            (
                "Pumping course (spinup → 3 yr pump)",
                "pumping_course_streamflow_storage.png",
                "Last 3 spinup years + 3 pumping years vs baseline.",
            ),
            (
                "WTD anomaly maps (pumping)",
                "pumping_wtd_anomaly_maps.png",
                "End spinup / end pump yr1 / end pump yr3.",
            ),
            (
                "Early vs buildup storage deficits",
                "pumping_early_buildup_storage_deficit_maps.png",
                "Early = end yr1 deficit; buildup = additional by end yr3.",
            ),
            (
                "Domain-total early vs buildup",
                "pumping_early_buildup_deficit_bars.png",
                "Map-based domain totals by pumping rate.",
            ),
        ],
    ),
]


def _set_run(run, text, size=28, bold=False, color=(0x1A, 0x1A, 0x1A)):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.font.name = "Calibri"


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(0.7), Inches(2.4), Inches(12), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, title, size=36, bold=True)
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    _set_run(run2, subtitle, size=18, color=(0x55, 0x55, 0x55))
    return slide


def add_section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(3.0), Inches(12), Inches(1.0))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, title, size=32, bold=True)
    return slide


def add_figure_slide(prs, title, image_path: Path | None, notes: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # title
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.55))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, title, size=22, bold=True)

    if image_path is not None and image_path.exists():
        # Fit inside content box preserving aspect ratio
        left, top = Inches(0.4), Inches(0.85)
        max_w, max_h = Inches(12.5), Inches(5.85)
        slide.shapes.add_picture(str(image_path), left, top, height=max_h)
        # If too wide, redo with width constraint
        pic = slide.shapes[-1]
        if pic.width > max_w:
            # remove and re-add with width
            sp = pic._element
            sp.getparent().remove(sp)
            slide.shapes.add_picture(str(image_path), left, top, width=max_w)
    else:
        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(12), Inches(4))
        tf = body.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, notes, size=18, color=(0x33, 0x33, 0x33))

    # footer notes when figure present
    if image_path is not None:
        foot = slide.shapes.add_textbox(Inches(0.4), Inches(6.9), Inches(12.5), Inches(0.45))
        tf = foot.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, notes, size=12, color=(0x66, 0x66, 0x66))
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # widescreen 16:9
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Drought recovery & pumping response figures",
        "Potomac & Wolf · 219 h products · August 2026",
    )

    missing = []
    for section_title, slides in SECTIONS:
        add_section_slide(prs, section_title)
        for title, fname, notes in slides:
            path = resolve_figure(fname) if fname else None
            if path is not None and not path.exists():
                missing.append(str(path))
                add_figure_slide(prs, title + " (missing)", None, f"Missing file: {path.name}\n\n{notes}")
            else:
                add_figure_slide(prs, title, path, notes)

    prs.save(OUT)
    print("wrote", OUT)
    if missing:
        print("missing figures:")
        for m in missing:
            print(" ", m)


if __name__ == "__main__":
    main()
